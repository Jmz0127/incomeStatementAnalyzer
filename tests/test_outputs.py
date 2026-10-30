from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from isviz import cli, narrative
from isviz.charts import overview
from isviz.export import RESTATED_FILL, to_excel
from isviz.statements import build_income_statement


@pytest.fixture
def stmt(facts):
    return build_income_statement(facts, "fix")


def test_excel_export(stmt, tmp_path):
    path = to_excel(stmt, tmp_path / "out.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Income Statement",
        "Margins",
        "YoY Growth",
        "Common Size",
        "XBRL Sources",
    ]
    ws = wb["Income Statement"]
    assert ws["B1"].value == "FY2019"
    labels = [c.value for c in ws["A"]]
    ni_row = labels.index("Net income") + 1
    fy2021_col = 4  # A=label, B=2019, C=2020, D=2021
    assert ws.cell(row=ni_row, column=fy2021_col).value == 110
    assert ws.cell(row=ni_row, column=fy2021_col).fill.fgColor.rgb.endswith(
        RESTATED_FILL.fgColor.rgb[-6:]
    )
    audit = list(wb["XBRL Sources"].values)
    assert ("Revenue", 2019) in {(r[0], r[1]) for r in audit[1:]}


def test_chart(stmt):
    fig = overview(stmt)
    assert len(fig.data) == 5
    assert fig.data[0].x[0] == "FY2019"


def test_cli_show_and_compare(stmt, monkeypatch):
    monkeypatch.setattr(cli, "load", lambda ticker, years: stmt)
    runner = CliRunner()
    result = runner.invoke(cli.app, ["show", "FIX"])
    assert result.exit_code == 0, result.output
    assert "Fixture Corp" in result.output
    assert "Revised between filings" in result.output
    result = runner.invoke(cli.app, ["compare", "FIX", "FIX"])
    assert result.exit_code == 0, result.output
    assert "gross_margin" in result.output


class FakeMessages:
    def __init__(self, response):
        self.response, self.kwargs = response, None

    def create(self, **kwargs):
        self.kwargs = kwargs
        return self.response


def _fake_client(monkeypatch, response):
    import anthropic

    messages = FakeMessages(response)
    client = SimpleNamespace(beta=SimpleNamespace(messages=messages))
    monkeypatch.setattr(anthropic, "Anthropic", lambda: client)
    return messages


def test_narrative_request_shape(stmt, monkeypatch):
    text = SimpleNamespace(type="text", text="Revenue grew 10.0% to $1.2 thousand.")
    messages = _fake_client(monkeypatch, SimpleNamespace(stop_reason="end_turn", content=[text]))
    assert narrative.generate(stmt).startswith("Revenue grew")
    kw = messages.kwargs
    assert kw["model"] == "claude-opus-5-5"
    assert kw["fallbacks"] == "default"
    assert "server-side-fallback-2026-07-01" in kw["betas"]
    assert '"restated_items"' in kw["messages"][0]["content"]


def test_narrative_refusal(stmt, monkeypatch):
    _fake_client(monkeypatch, SimpleNamespace(stop_reason="refusal", content=[]))
    with pytest.raises(narrative.NarrativeError, match="declined"):
        narrative.generate(stmt)


def test_mcp_tools_registered():
    import asyncio

    from isviz.mcp_server import server

    names = {t.name for t in asyncio.run(server.list_tools())}
    assert names == {"get_income_statement", "get_financial_summary", "compare_companies"}
