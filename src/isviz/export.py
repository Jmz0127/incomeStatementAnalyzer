"""Review-ready Excel workbook: statement, ratios, and an XBRL audit trail."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from isviz import analysis
from isviz.statements import LABELS, IncomeStatement

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RESTATED_FILL = PatternFill("solid", fgColor="FFF2CC")
MILLIONS = '#,##0.0,,;(#,##0.0,,);"–"'
PERCENT = '0.0%;(0.0%);"–"'
EPS = '0.00;(0.00);"–"'


def _style(ws: Worksheet, number_format: str, first_col_width: int = 30) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = first_col_width
    for col in range(2, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12
    for row in ws.iter_rows(min_row=2, min_col=2):
        for cell in row:
            cell.number_format = number_format
    ws.freeze_panes = "B2"


def to_excel(stmt: IncomeStatement, path: str | Path) -> Path:
    path = Path(path)
    years = [int(y) for y in stmt.data.index]
    statement = stmt.labeled()
    statement.columns = [f"FY{y}" for y in years]

    sheets: dict[str, tuple[pd.DataFrame, str]] = {
        "Income Statement": (statement, MILLIONS),
        "Margins": (
            analysis.margins(stmt).T.rename(index=lambda s: s.replace("_", " ").capitalize()),
            PERCENT,
        ),
        "YoY Growth": (analysis.yoy_growth(stmt).T.rename(index=LABELS), PERCENT),
        "Common Size": (analysis.common_size(stmt).T.rename(index=LABELS), PERCENT),
    }

    audit_rows = [
        {
            "line_item": LABELS[item],
            "fiscal_year": fy,
            "period_end": stmt.period_end.get(fy),
            "xbrl_concept": concept,
            "restated": fy in stmt.restated.get(item, []),
        }
        for item, by_year in stmt.sources.items()
        for fy, concept in sorted(by_year.items())
    ]

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, (df, fmt) in sheets.items():
            df.columns = [f"FY{y}" if isinstance(y, int) else y for y in df.columns]
            df.to_excel(xl, sheet_name=name)
            _style(xl.sheets[name], fmt)
        pd.DataFrame(audit_rows).to_excel(xl, sheet_name="XBRL Sources", index=False)

        ws = xl.sheets["Income Statement"]
        ws.cell(row=1, column=1, value=f"{stmt.company} ($ millions)")
        eps_row = list(statement.index).index(LABELS["eps_diluted"]) + 2
        for cell in ws[eps_row][1:]:
            cell.number_format = EPS
        # Highlight values that changed between filings so a reviewer can tie them out.
        row_of = {item: list(statement.index).index(LABELS[item]) + 2 for item in stmt.restated}
        for item, fys in stmt.restated.items():
            for fy in fys:
                ws.cell(row=row_of[item], column=years.index(fy) + 2).fill = RESTATED_FILL

        audit = xl.sheets["XBRL Sources"]
        for cell in audit[1]:
            cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        for col, width in zip("ABCDE", (28, 12, 12, 70, 10), strict=True):
            audit.column_dimensions[col].width = width
    return path
